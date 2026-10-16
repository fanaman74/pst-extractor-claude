import io

from app.extractors import decode_bytes, extract_attachment_text, legacy_binary_text, safe_filename


def test_safe_filename_blocks_traversal():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("..\\..\\boot.ini") == "boot.ini"
    assert safe_filename("") == "attachment"
    assert safe_filename("rëport <final>.pdf") == "r_port _final_.pdf"


def test_decode_bytes_fallbacks():
    assert decode_bytes("héllo".encode("utf-8")) == "héllo"
    assert decode_bytes("héllo".encode("cp1252")) == "héllo"
    assert decode_bytes("hi".encode("utf-16")) == "hi"
    assert decode_bytes("Grüße".encode("cp1252"), codepage=1252) == "Grüße"


def test_text_csv_html():
    assert extract_attachment_text("a.txt", b"plain text")[0] == "plain text"
    assert extract_attachment_text("a.csv", b"a,b\n1,2\n")[0] == "a\tb\n1\t2"
    text, status = extract_attachment_text("a.html", b"<html><script>x()</script><p>Hello <b>there</b></p></html>")
    assert "Hello" in text and "x()" not in text and status == "ok"


def test_unsupported_and_empty():
    assert extract_attachment_text("photo.jpg", b"\xff\xd8\xff") == ("", "unsupported")
    assert extract_attachment_text("blank.txt", b"   ")[1] == "empty"


def test_docx_xlsx_pdf_roundtrip():
    import docx
    import openpyxl

    d = docx.Document()
    d.add_paragraph("Quarterly pipeline capacity report")
    buf = io.BytesIO(); d.save(buf)
    assert "pipeline capacity" in extract_attachment_text("r.docx", buf.getvalue())[0]

    wb = openpyxl.Workbook(); wb.active.append(["Counterparty", "Volume"]); wb.active.append(["Citizens", 1200])
    buf = io.BytesIO(); wb.save(buf)
    text = extract_attachment_text("s.xlsx", buf.getvalue())[0]
    assert "Citizens\t1200" in text

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    pdf = fitz.open(); page = pdf.new_page(); page.insert_text((72, 72), "Imbalance statement for Transwestern")
    assert "Transwestern" in extract_attachment_text("x.pdf", pdf.tobytes())[0]


def test_corrupt_file_reports_error_instead_of_raising():
    text, status = extract_attachment_text("broken.docx", b"not a zip")
    assert text == "" and status.startswith("error:")


def test_legacy_doc_best_effort():
    blob = b"\x00\x01binary\x00" + b"This agreement is made between the parties\r" + b"\x07" + "Unicode text inside the document".encode("utf-16-le")
    out = legacy_binary_text(blob)
    assert "This agreement is made between the parties" in out
    assert "Unicode text inside the document" in out
