"""Text extraction for message bodies and attachments. Every function is defensive: bad input
returns ('', status) instead of raising, so one broken file never fails an archive."""
from __future__ import annotations

import csv
import io
import os
import re
import warnings

TEXT_EXTS = {".txt", ".csv", ".tsv", ".md", ".json", ".xml", ".log", ".eml", ".ics", ".vcf", ".url", ".ini", ".yaml", ".yml"}
HTML_EXTS = {".htm", ".html"}


def decode_bytes(data: bytes | str | None, codepage: int | None = None) -> str:
    if data is None:
        return ""
    if isinstance(data, str):
        return data
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    candidates = []
    if codepage:
        candidates.append("utf-8" if codepage == 65001 else f"cp{codepage}")
    candidates += ["utf-8", "cp1252"]
    for enc in candidates:
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1", errors="replace")


def html_to_text(html: str) -> str:
    from bs4 import BeautifulSoup

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "head", "title", "meta"]):
        tag.decompose()
    return soup.get_text("\n")


def rtf_to_text(rtf: str) -> str:
    from striprtf.striprtf import rtf_to_text as _strip

    try:
        return _strip(rtf, errors="ignore")
    except Exception:
        return ""


_PRINTABLE_RUN = re.compile(rb"[\x20-\x7e\t\r\n]{6,}")
_UTF16_RUN = re.compile(rb"(?:[\x20-\x7e]\x00){6,}")


def legacy_binary_text(data: bytes) -> str:
    """Best effort for legacy Office binaries (.doc/.ppt): printable ASCII and UTF-16LE runs."""
    # Word uses \r for paragraphs, \x07 for table cells and \x0b/\x0c for line/page breaks
    data = data.translate(bytes.maketrans(b"\x07\x0b\x0c", b"\n\n\n"))
    runs = [m.group().decode("ascii", "ignore") for m in _PRINTABLE_RUN.finditer(data)]
    runs += [m.group().decode("utf-16-le", "ignore") for m in _UTF16_RUN.finditer(data)]
    # Drop short noise and obvious binary/format artefacts
    keep = [r.strip() for r in runs if len(r.strip()) >= 20 and sum(c.isalpha() for c in r) > len(r) * 0.5]
    return "\n".join(keep)


def _pdf(data: bytes) -> str:
    try:
        import pymupdf as fitz
    except ImportError:  # PyMuPDF < 1.24.3
        import fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes, max_chars: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, n = [], 0
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            line = "\t".join("" if v is None else str(v) for v in row).strip()
            if line:
                out.append(line)
                n += len(line)
                if n > max_chars:
                    return "\n".join(out)
    return "\n".join(out)


def _xls(data: bytes, max_chars: int) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    out, n = [], 0
    for sheet in book.sheets():
        out.append(f"# Sheet: {sheet.name}")
        for r in range(sheet.nrows):
            line = "\t".join(str(v) for v in sheet.row_values(r) if v not in ("", None)).strip()
            if line:
                out.append(line)
                n += len(line)
                if n > max_chars:
                    return "\n".join(out)
    return "\n".join(out)


def _csv(data: bytes) -> str:
    text = decode_bytes(data)
    try:
        return "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text)))
    except csv.Error:
        return text


def extract_attachment_text(filename: str, data: bytes, max_chars: int = 300_000) -> tuple[str, str]:
    """Returns (text, status). status is one of: ok, best_effort, empty, unsupported, error:<msg>."""
    ext = os.path.splitext(filename.lower())[1]
    try:
        if ext == ".pdf" or data[:5] == b"%PDF-":
            text, status = _pdf(data), "ok"
        elif ext in (".docx", ".docm"):
            text, status = _docx(data), "ok"
        elif ext in (".xlsx", ".xlsm"):
            text, status = _xlsx(data, max_chars), "ok"
        elif ext == ".xls":
            text, status = _xls(data, max_chars), "ok"
        elif ext in (".csv", ".tsv"):
            text, status = _csv(data), "ok"
        elif ext in HTML_EXTS:
            text, status = html_to_text(decode_bytes(data)), "ok"
        elif ext == ".rtf":
            text, status = rtf_to_text(decode_bytes(data)), "ok"
        elif ext in TEXT_EXTS:
            text, status = decode_bytes(data), "ok"
        elif ext in (".doc", ".ppt", ".dot"):
            text, status = legacy_binary_text(data), "best_effort"
        else:
            return "", "unsupported"
    except Exception as exc:  # corrupt or password-protected files
        return "", f"error:{type(exc).__name__}: {exc}"[:300]
    text = text.replace("\x00", "")[:max_chars]
    return (text, status) if text.strip() else ("", "empty")


_UNSAFE = re.compile(r"[^A-Za-z0-9._ ()+-]+")


def safe_filename(name: str | None, fallback: str = "attachment") -> str:
    name = os.path.basename((name or "").replace("\\", "/")).strip()
    name = _UNSAFE.sub("_", name).strip(" .")
    return (name or fallback)[:150]
